"""偏差报告 API"""

import logging
import uuid
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, UploadFile
from fastapi.responses import StreamingResponse

from app.core import storage
from app.core.deps import PermissionContext, RequiredUser, require_permission
from app.core.exceptions import AppException, NotFoundException
from app.core.llm import llm_client
from app.core.response import DataResponse
from app.modules.quality.qms.deviation_report_schemas import (
    DeviationOptimizeData,
    DeviationReportUploadData,
    OptimizeTextRequest,
    SaveHtmlRequest,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/deviation-report", tags=["偏差报告"])

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# 原始文档存 MinIO 的 quality 模块桶，容器重建不丢；key 前缀与业务目录一致
_STORAGE_MODULE = "quality"
_STORAGE_PREFIX = "deviation_reports"

# 优化类型 → 提示词模板
_OPTIMIZE_PROMPTS: dict[str, str] = {
    "polish": """请对以下文本进行专业润色，保持原意，使语言更加流畅、专业、规范。

原文：
{text}

要求：
1. 保持原文的核心意思不变
2. 修正语法错误和表达不当之处
3. 使语言更加流畅、专业
4. 保持原有的格式和结构
5. 直接返回润色后的文本，不要添加任何解释说明

润色结果：""",
    "summarize": """请对以下文本进行简要总结。

原文：
{text}

要求：
1. 提取核心要点
2. 语言简洁明了
3. 直接返回总结内容，不要添加任何解释说明

总结结果：""",
    "expand": """请对以下文本进行适当扩展和详细说明。

原文：
{text}

要求：
1. 在保持原意的基础上进行扩展
2. 增加必要的细节和说明
3. 使内容更加完整和详尽
4. 直接返回扩展后的文本，不要添加任何解释说明

扩展结果：""",
    "simplify": """请对以下文本进行简化，使其更加简洁易懂。

原文：
{text}

要求：
1. 去除冗余表达
2. 使用简洁清晰的语言
3. 保持核心意思不变
4. 直接返回简化后的文本，不要添加任何解释说明

简化结果：""",
}

_OPTIMIZE_RESULT_PREFIXES = ("润色结果：", "总结结果：", "扩展结果：", "简化结果：")


def _object_key(file_id: str) -> str:
    return f"{_STORAGE_PREFIX}/{file_id}"


def _store_bytes(file_id: str, data: bytes, content_type: str) -> None:
    """把原始文档写入对象存储。"""
    if not storage.is_enabled():
        raise AppException(status_code=503, message="对象存储未启用，无法保存原始文档")
    storage.upload_object(_STORAGE_MODULE, _object_key(file_id), data, len(data), content_type)


def _load_bytes(file_id: str) -> bytes:
    """按 file_id 读取原始文档，拒绝任何目录穿越写法。"""
    if not file_id or Path(file_id).name != file_id:
        raise AppException(status_code=400, message="文件标识不合法")
    found = storage.get_object(_STORAGE_MODULE, _object_key(file_id))
    if found is None:
        raise NotFoundException(resource="原始文件")
    return found[0]


def _docx_response(data: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        BytesIO(data),
        media_type=DOCX_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ============ Word 模板上传 API ============


@router.post("/upload", summary="上传并解析Word模板")
async def upload_report(
    current_user: RequiredUser,
    ctx: PermissionContext = Depends(require_permission("quality:deviation:update")),
    file: UploadFile = File(...),
) -> DataResponse[DeviationReportUploadData]:
    """上传Word文档并解析为HTML格式"""
    import mammoth  # type: ignore[import-untyped]

    if not file.filename or not file.filename.lower().endswith((".docx", ".doc")):
        raise AppException(status_code=400, message="仅支持 .docx 或 .doc 格式的Word文件")

    file_ext = Path(file.filename).suffix.lower()
    saved_filename = f"{uuid.uuid4().hex}{file_ext}"
    raw = await file.read()
    _store_bytes(saved_filename, raw, file.content_type or DOCX_MEDIA_TYPE)

    try:
        result = mammoth.convert_to_html(BytesIO(raw))
        html_content = result.value
        warnings = [msg.message for msg in result.messages]
    except Exception:
        logger.exception("偏差报告解析失败: file_id=%s", saved_filename)
        try:
            storage.delete_object(_STORAGE_MODULE, _object_key(saved_filename))
        except Exception:
            logger.exception("偏差报告上传清理失败: file_id=%s", saved_filename)
        raise AppException(status_code=500, message="文档解析失败，请确认文件是有效的 Word 文档") from None

    return DataResponse[DeviationReportUploadData](
        message="文档解析成功",
        data=DeviationReportUploadData(
            file_id=saved_filename,
            file_name=file.filename,
            html_content=html_content,
            file_path=f"/api/v1/quality/deviation-report/download?file_id={saved_filename}",
            warnings=warnings,
        ),
    )


@router.post("/save", summary="保存HTML内容为Word")
async def save_report(
    payload: SaveHtmlRequest,
    current_user: RequiredUser,
    ctx: PermissionContext = Depends(require_permission("quality:deviation:update")),
) -> StreamingResponse:
    """将HTML内容保存为Word文档"""
    from bs4 import BeautifulSoup
    from docx import Document

    doc = Document(BytesIO(_load_bytes(payload.file_id)))
    soup = BeautifulSoup(payload.html_content, "html.parser")

    # 清除所有现有段落（保留样式）
    for para in doc.paragraphs:
        para.clear()  # type: ignore[no-untyped-call]

    for element in soup.children:
        tag = getattr(element, "name", None)
        if tag == "p":
            para = doc.add_paragraph()
            for child in getattr(element, "children", []):
                child_tag = getattr(child, "name", None)
                if child_tag:
                    if child_tag in ("b", "strong"):
                        run = para.add_run(child.get_text())
                        run.bold = True
                    elif child_tag == "i":
                        run = para.add_run(child.get_text())
                        run.italic = True
                    else:
                        para.add_run(child.get_text())
                else:
                    para.add_run(str(child))
        elif tag == "table":
            # 表格保持原样，跳过重建
            continue
        elif hasattr(element, "get_text") and element.get_text().strip():
            para = doc.add_paragraph()
            para.add_run(element.get_text())

    buf = BytesIO()
    doc.save(buf)
    return _docx_response(buf.getvalue(), f"偏差报告_{uuid.uuid4().hex[:8]}.docx")


@router.get("/download", summary="下载原始Word文档")
async def download_report(
    current_user: RequiredUser,
    ctx: PermissionContext = Depends(require_permission("quality:deviation:read")),
    file_id: str = "",
) -> StreamingResponse:
    """下载原始Word文档"""
    return _docx_response(_load_bytes(file_id), file_id)


# ============ AI 优化接口 ============


@router.post("/ai/optimize", summary="AI优化文本")
async def optimize_text(
    payload: OptimizeTextRequest,
    current_user: RequiredUser,
    ctx: PermissionContext = Depends(require_permission("quality:deviation:update")),
) -> DataResponse[DeviationOptimizeData]:
    """使用AI优化选中的文本"""
    template = _OPTIMIZE_PROMPTS.get(payload.optimize_type, _OPTIMIZE_PROMPTS["polish"])
    prompt = template.format(text=payload.text)

    try:
        result = await llm_client.chat(
            [{"role": "user", "content": prompt}],
            response_format=None,
            max_tokens=4096,
        )
    except Exception:
        logger.exception("偏差报告 AI 优化失败: optimize_type=%s", payload.optimize_type)
        raise AppException(status_code=500, message="AI 优化失败，请稍后重试") from None

    optimized = result.strip()
    for prefix in _OPTIMIZE_RESULT_PREFIXES:
        if optimized.startswith(prefix):
            optimized = optimized[len(prefix) :].lstrip()
            break

    return DataResponse[DeviationOptimizeData](
        message="优化成功",
        data=DeviationOptimizeData(optimized_text=optimized),
    )
