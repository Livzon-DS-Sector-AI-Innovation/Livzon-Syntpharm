"""Document parser for knowledge base import."""

import io
import os
import re
from typing import Any

# 分类关键词映射
CATEGORY_KEYWORDS = {
    "laws_regulations": [
        "法",
        "条例",
        "规定",
        "办法",
        "安全生产法",
        "职业病防治法",
        "消防法",
        "环境保护法",
        "劳动法",
        "合同法",
    ],
    "standards": ["GB", "GB/T", "标准", "规范", "ISO", "AQ", "HG", "SH", "国家标准", "行业标准", "地方标准"],
    "management_systems": ["制度", "规程", "管理办法", "责任制", "操作规范", "作业指导书", "管理程序", "工作流程"],
    "accident_cases": ["事故", "案例", "通报", "调查报告", "事故分析", "事故案例"],
    "emergency_plans": ["预案", "应急", "处置方案", "应急预案", "应急响应", "应急救援"],
    "sds": ["SDS", "MSDS", "化学品安全技术说明书", "安全数据表", "物质安全数据表"],
    "training_materials": ["培训", "教材", "课件", "教案", "讲义", "培训资料", "考试题库"],
}


def infer_category(filename: str) -> str:
    """根据文件名推断分类"""
    name_lower = filename.lower()

    for category, keywords in CATEGORY_KEYWORDS.items():
        for keyword in keywords:
            if keyword.lower() in name_lower:
                return category

    return "other"


def extract_title(filename: str) -> str:
    """从文件名提取标题"""
    name = os.path.splitext(filename)[0]

    # 去掉常见后缀（版本号、日期等）
    name = re.sub(r"[_\-]\d{4}.*$", "", name)
    name = re.sub(r"[_\-]v\d+.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"[_\-]版本.*$", "", name)

    name = name.strip()
    name = re.sub(r"[_\-]+", " ", name)

    return name if name else filename


def extract_tags(filename: str) -> str | None:
    """从文件名提取标签"""
    tags = []
    name = filename.lower()

    if "2021" in name or "2021修订" in name:
        tags.append("2021版")
    if "修订" in name:
        tags.append("修订版")
    if "最新" in name:
        tags.append("最新")

    return ",".join(tags) if tags else None


def parse_pdf(content: bytes) -> str:
    """解析 PDF 文件"""
    try:
        import pdfplumber

        text_parts = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for i, page in enumerate(pdf.pages[:10]):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

        return "\n\n".join(text_parts)
    except Exception as e:
        return f"[PDF 解析失败: {str(e)}]"


def parse_docx(content: bytes) -> str:
    """解析 Word 文档"""
    try:
        from docx import Document

        doc = Document(io.BytesIO(content))
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        return "\n\n".join(text_parts)
    except Exception as e:
        return f"[Word 解析失败: {str(e)}]"


def parse_excel(content: bytes) -> str:
    """解析 Excel 文件"""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True)
        text_parts = []

        for sheet_name in wb.sheetnames[:3]:
            ws = wb[sheet_name]
            text_parts.append(f"## {sheet_name}\n")

            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if row_count >= 50:
                    text_parts.append("... (更多数据省略)")
                    break

                row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                if row_text.strip(" |"):
                    text_parts.append(row_text)
                    row_count += 1

        return "\n\n".join(text_parts)
    except Exception as e:
        return f"[Excel 解析失败: {str(e)}]"


def parse_txt(content: bytes) -> str:
    """解析纯文本文件"""
    try:
        return content.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[文本解析失败: {str(e)}]"


async def parse_document(filename: str, content: bytes) -> dict[str, Any]:
    """
    解析文档，提取标题、摘要、正文、分类等信息

    Args:
        filename: 文件名
        content: 文件内容（字节）

    Returns:
        dict: {
            "title": str,
            "summary": str | None,
            "content": str | None,
            "category": str,
            "tags": str | None,
        }
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        text = parse_pdf(content)
    elif ext == ".docx":
        text = parse_docx(content)
    elif ext in [".xlsx", ".xls"]:
        text = parse_excel(content)
    elif ext in [".txt", ".md"]:
        text = parse_txt(content)
    else:
        text = f"[不支持的文件格式: {ext}]"

    title = extract_title(filename)
    category = infer_category(filename)
    tags = extract_tags(filename)

    # 生成摘要（取前 500 字符）
    summary = text[:500].strip() if text and not text.startswith("[") else None

    return {
        "title": title,
        "summary": summary,
        "content": text if text and not text.startswith("[") else None,
        "category": category,
        "tags": tags,
    }
