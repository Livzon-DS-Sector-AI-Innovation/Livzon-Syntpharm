"""docx 母本 → 全内容 Markdown（基于已有 python-docx 依赖，零新增包）。

定位：交付物模板页「模板 Markdown」弹窗的数据源——把 Word 模板文档整体解析为
Markdown 全文（标题、段落、列表、表格按文档流顺序输出），而不是填写项骨架。

设计约束：
- 按 ``body`` 子元素顺序遍历，段落与表格交错不错位（python-docx 的
  ``doc.paragraphs``/``doc.tables`` 分离迭代做不到这一点）。
- 标题识别只认样式（``Heading 1``–``Heading 6`` / 中文 ``标题 1``–``标题 6`` /
  ``Title``），不做「加粗短行猜标题」的启发式，避免把强调文字误判成章节。
- 格式保真：run 级加粗/斜体转 ``**``/``*``（只认显式格式，样式继承的 None 不猜）；
  列表按编号层级 ``ilvl`` 缩进，多级列表不再被压平。
- 表格输出 GFM：单元格内换行转 ``<br>``、竖线转义；合并单元格按去重后的实际格
  输出（``row.cells`` 会把横向合并重复展开成多列同文本，原样输出会列错位、
  文本重复），列数按全表最大值右侧补空；首行全空时补「列N」表头占位，
  保证 GFM 分隔行有对应表头；**只有单列的表**（标题横幅、竖排标签）改按段落输出。
- 目录条目（``TOC``/``目录`` 样式，或行尾「点线/制表符 + 页码」的域结果）跳过：
  在 Markdown 里既点不动，又会把正文刷成一屏噪音。
- 页眉页脚、图片、批注、目录域不进入文本视图。

只读转换、不落库；成文链路（锚点回填母本 docx）与本文件无关。
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.modules.research.doc_gen.grid import distinct_row_cells

try:  # python-docx >= 1.1 才有 Hyperlink；低版本回退 runs 拼接
    from docx.text.hyperlink import Hyperlink
except ImportError:  # pragma: no cover
    Hyperlink = None  # type: ignore[assignment,misc]

# 样式名 → 标题级别：兼容英文 Word 与中文版式（"标题 1"/"标题1"）
_HEADING_RE = re.compile(r"^(?:heading|标题)\s*([1-6])\s*$", re.IGNORECASE)
_TITLE_STYLES = {"title", "标题"}
# 列表样式（除 numPr 外的兜底识别，如 "List Paragraph" 系）
_LIST_STYLE_RE = re.compile(r"^(list|列表)", re.IGNORECASE)
# 目录条目样式（"TOC 1"/"目录 1"）：整行是「标题文本 + 制表符 + 页码」的域结果，
# 在 Markdown 视图里既不能跳转又刷屏，且会被当成正文噪音，直接跳过
_TOC_STYLE_RE = re.compile(r"^(toc|目录)", re.IGNORECASE)
# 兜底：样式名不规范的目录行（行尾「制表符，或 ≥3 个点线字符 + 页码」）。
# 要求制表符或连续点线，单个空格加数字（「共 12」）不算，避免误杀正文。
_TOC_PAGENUM_RE = re.compile(r"(?:\t|[.\u00b7\u2026 ]{3,})\s*\d{1,4}$")

_MULTI_BLANK_RE = re.compile(r"\n{3,}")

# 正文文本里会被误当 Markdown 语法的字符：反斜杠、反引号、星号/下划线（强调）、
# 尖括号（HTML/自动链接，rehype-raw 下还有注入风险）、方括号（链接）。母本正文常见
# 涂黑公司名 ``****``、``*100%`` 等，不转义就会被当成强调符——星号被吃掉并错加粗斜体。
_MD_ESCAPE_RE = re.compile(r"([\\`*_<>[\]])")


def _escape_md_text(text: str) -> str:
    """转义文本 run 里的 Markdown 语法字符；结构标记（# 标题、- 列表、| 表格、<br>）由调用方在转义之后再加。"""
    return _MD_ESCAPE_RE.sub(r"\\\1", text)


def _heading_level(style_name: str | None) -> int:
    """样式名映射标题级别；非标题样式返回 0。"""
    if not style_name:
        return 0
    name = style_name.strip()
    match = _HEADING_RE.match(name)
    if match:
        return int(match.group(1))
    if name.lower() in _TITLE_STYLES:
        return 1
    return 0


def _is_list_item(para: Paragraph) -> bool:
    """项目符号/编号列表：优先看 numPr，其次看列表样式名。"""
    ppr = para._p.pPr
    if ppr is not None and ppr.numPr is not None:
        return True
    style_name = para.style.name if para.style is not None else None
    return bool(style_name and _LIST_STYLE_RE.match(style_name.strip()))


def _run_md(run: Run) -> str:
    """单个 run → Markdown：先转义再加粗/斜体标记；首尾空白留在标记外避免渲染歧义。

    只认显式格式（``bold``/``italic`` 为 True）；``None`` 表示样式继承，不猜，
    避免把整段正文都错标成强调。
    """
    text = _escape_md_text(run.text)
    if not text.strip():
        return text
    emphasis = "***" if (run.bold and run.italic) else ("**" if run.bold else "*" if run.italic else "")
    if not emphasis:
        return text
    lead = text[: len(text) - len(text.lstrip())]
    tail = text[len(text.rstrip()) :]
    return f"{lead}{emphasis}{text.strip()}{emphasis}{tail}"


def _inline_md(para: Paragraph) -> str:
    """段落内联内容 → Markdown 文本（run 保真；超链接取文本；老版本回退 para.text）。"""
    if Hyperlink is not None and hasattr(para, "iter_inner_content"):
        parts = [
            _run_md(item) if isinstance(item, Run) else _escape_md_text(item.text)
            for item in para.iter_inner_content()
            if isinstance(item, (Run, Hyperlink))
        ]
        text = "".join(parts)
        if text.strip() or not para.text.strip():
            return text
    return _escape_md_text(para.text)


def _list_prefix(para: Paragraph) -> str:
    """列表项前缀：编号层级 ilvl → 每级两个空格，多级列表不再被压平。"""
    if not _is_list_item(para):
        return ""
    ilvl = 0
    ppr = para._p.pPr
    # numPr/ilvl 是 python-docx 动态 oxml 元素，类型桩无属性声明，用 getattr 取
    num_pr = getattr(ppr, "numPr", None) if ppr is not None else None
    ilvl_el = getattr(num_pr, "ilvl", None) if num_pr is not None else None
    if ilvl_el is not None:
        try:
            ilvl = max(0, int(ilvl_el.val or 0))
        except (TypeError, ValueError):
            ilvl = 0
    return f"{'  ' * ilvl}- "


def _is_toc_entry(style_name: str | None, text: str) -> bool:
    """目录条目：样式名命中，或「正文 + 点线/制表符 + 页码」的域结果形态。"""
    if style_name and _TOC_STYLE_RE.match(style_name.strip()):
        return True
    return bool(_TOC_PAGENUM_RE.search(text))


def _paragraph_lines(para: Paragraph) -> list[str]:
    text = _inline_md(para).strip()
    if not text:
        return []
    style_name = para.style.name if para.style is not None else None
    level = _heading_level(style_name)
    # 目录判定必须用原文：制表符压成空格后就再也认不出「标题 + 页码」形态了
    if not level and _is_toc_entry(style_name, text):
        return []
    text = text.replace("\t", " ")
    if level:
        return ["", f"{'#' * level} {text}", ""]
    return [f"{_list_prefix(para)}{text}", ""]


def _cell_md(cell: _Cell) -> str:
    """单元格 → 一行表格内容：段落/软换行转 ``<br>``、run 保真、竖线转义防撑破 GFM。"""
    parts = [re.sub(r"\s*\n\s*", "<br>", _inline_md(p).strip()) for p in cell.paragraphs]
    text = "<br>".join(part for part in parts if part)
    return text.replace("|", "\\|")


def _table_lines(table: Table) -> list[str]:
    rows = list(table.rows)
    if not rows:
        return []
    # 合并单元格去重：row.cells 把横向合并重复展开，直接输出会列错位+文本重复
    grid = [[_cell_md(cell) for cell in distinct_row_cells(row)] for row in rows]
    width = max(len(cells) for cells in grid)
    if width == 0:
        return []
    if width == 1:
        # 单列表格几乎全是「标题横幅 / 竖排标签」，输出成只有一列的 GFM 表既难看又
        # 抢不走语义，按段落输出更贴近原文观感
        paragraphs: list[str] = []
        for cells in grid:
            text = cells[0].strip()
            if text:
                paragraphs.extend(["", text, ""])
        return paragraphs
    if not any(cell for cell in grid[0]):
        grid[0] = [f"列{pos + 1}" for pos in range(width)]  # GFM 分隔行需要表头，全空首行补占位
    lines: list[str] = [""]
    for index, cells in enumerate(grid):
        cells += [""] * (width - len(cells))
        lines.append("| " + " | ".join(cells) + " |")
        if index == 0:
            lines.append("| " + " | ".join(["---"] * width) + " |")
    lines.append("")
    return lines


def render_docx_markdown(data: bytes) -> str:
    """把 docx 字节流转为全内容 Markdown 文本。

    损坏/非法 docx 会抛 ``docx`` 底层异常（zipfile.BadZipFile、ValueError 等），
    由调用方决定回退策略（本模块不做吞异常的兜底）。
    """
    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            lines.extend(_paragraph_lines(Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            lines.extend(_table_lines(Table(child, doc)))
    text = _MULTI_BLANK_RE.sub("\n\n", "\n".join(lines)).strip()
    return f"{text}\n" if text else ""


__all__ = ["render_docx_markdown"]
