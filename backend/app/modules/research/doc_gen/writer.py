"""docx 文本写入原语：只改文字，不改样式。

所有写入都复用目标位置（或同行标签单元格）的第一个 run 属性，从而保证字体、字号、
加粗、颜色与母本一致；表格扩行通过 deepcopy 模板行 XML 实现，边框/行高/列宽随模板。
"""

from __future__ import annotations

import re
from copy import deepcopy
from typing import TYPE_CHECKING, cast

from docx.table import Table, _Cell, _Row
from docx.text.paragraph import Paragraph

from app.modules.research.doc_gen.anchors import analyze_data_area
from app.modules.research.doc_gen.grid import cell_at_grid, distinct_row_cells

if TYPE_CHECKING:
    from app.modules.research.doc_gen.anchors import RowsTarget
    from app.modules.research.doc_gen.template_spec import TableSpec

_COLONS = ("：", ":")
_WS = re.compile(r"\s+")


def _first_run_properties(paragraph: Paragraph) -> object | None:
    """取段落第一个 run 的 rPr（格式），无则 None。"""
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return run._r.rPr
    return None


def _clear_runs(paragraph: Paragraph) -> None:
    """删除段落内所有 run（保留段落属性与样式）。"""
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)


def set_paragraph_text(paragraph: Paragraph, text: str, style_from: Paragraph | None = None) -> None:
    """整段替换文本，沿用原段落第一个 run 的格式。"""
    props = _first_run_properties(paragraph)
    if props is None and style_from is not None:
        props = _first_run_properties(style_from)
    _clear_runs(paragraph)
    run = paragraph.add_run(text or "")
    if props is not None:
        run._r.insert(0, deepcopy(props))


def split_label(text: str) -> tuple[str, str]:
    """把「标签：值」拆成标签部分（含冒号）与值部分；无冒号时标签为空。"""
    for colon in _COLONS:
        if colon in text:
            head, _, tail = text.partition(colon)
            return f"{head}{colon}", tail
    return "", text


def set_after_label(paragraph: Paragraph, text: str) -> None:
    """保留「标签：」前缀，替换冒号之后的内容。"""
    label, _old = split_label(paragraph.text)
    value = text or ""
    if not label:
        set_paragraph_text(paragraph, value)
        return
    set_paragraph_text(paragraph, f"{label}{value}")


def write_multiline(
    paragraph: Paragraph, text: str, mode: str = "replace", label: str = "", style_from: Paragraph | None = None
) -> list[Paragraph]:
    """写入可能多行的文本；额外行以克隆段落插入其后（样式与母本段落一致）。"""
    lines = [line for line in (text or "").splitlines() if line.strip() != ""] or [""]
    if mode == "after_label":
        head, _ = split_label(paragraph.text)
        lines[0] = f"{head}{lines[0]}"
    elif mode == "append_after":
        pass
    set_paragraph_text(paragraph, lines[0], style_from=style_from)
    written = [paragraph]
    anchor = paragraph
    for line in lines[1:]:
        clone = deepcopy(anchor._p)
        anchor._p.addnext(clone)
        new_para = Paragraph(clone, anchor._parent)
        set_paragraph_text(new_para, line, style_from=paragraph)
        written.append(new_para)
        anchor = new_para
    return written


def cell_paragraphs(cell: _Cell) -> list[Paragraph]:
    """单元格段落列表（显式类型，避免 Any 扩散）。"""
    return list(cell.paragraphs)


def cell_style_source(cell: _Cell) -> Paragraph | None:
    """从单元格中找一个带格式的段落作为格式来源。"""
    for para in cell_paragraphs(cell):
        if _first_run_properties(para) is not None:
            return para
    return None


def set_cell_text(cell: _Cell, text: str, style_from: _Cell | None = None) -> None:
    """写单元格：优先复用本格格式，其次复用来源格格式。"""
    paras = cell_paragraphs(cell)
    target = paras[0] if paras else cast(Paragraph, cell.add_paragraph())
    source = cell_style_source(cell)
    if source is None and style_from is not None:
        source = cell_style_source(style_from)
    lines = [line for line in (text or "").splitlines() if line.strip() != ""] or [""]
    set_paragraph_text(target, lines[0], style_from=source)
    anchor = target
    for line in lines[1:]:
        clone = deepcopy(anchor._p)
        anchor._p.addnext(clone)
        new_para = Paragraph(clone, anchor._parent)
        set_paragraph_text(new_para, line, style_from=source)
        anchor = new_para


def row_cells(table: Table, row_index: int) -> list[_Cell]:
    """行内去重单元格。"""
    return distinct_row_cells(table.rows[row_index])


def write_cell(table: Table, row_index: int, grid_col: int, text: str) -> None:
    """按网格列写单元格，格式来源取同行第一个单元格（通常是标签格）。"""
    cells = row_cells(table, row_index)
    target = cell_at_grid(table, row_index, grid_col)
    if target is None:
        return
    set_cell_text(target, text, style_from=cells[0] if cells else None)


def replace_keyword(paragraph: Paragraph, keyword: str, value: str) -> int:
    """替换段落中出现的占位关键词（如 XXXX），返回替换次数。

    逐 run 替换以保留格式；关键词被拆分在多个 run 时退化为整段重写。
    """
    if keyword not in paragraph.text:
        return 0
    count = 0
    for run in paragraph.runs:
        if keyword in run.text:
            occurrences = run.text.count(keyword)
            run.text = run.text.replace(keyword, value)
            count += occurrences
    if count == 0:
        set_paragraph_text(paragraph, paragraph.text.replace(keyword, value))
        count = paragraph.text.count(value)
    return count


def _row_text(row: _Row) -> str:
    """行文本（判断是否空行用）。"""
    return _WS.sub("", "".join(cell.text for cell in distinct_row_cells(row)))


def _clear_row(table: Table, row_index: int) -> None:
    """清空一行所有单元格文字（保留格式）。"""
    for cell in row_cells(table, row_index):
        for para in cell_paragraphs(cell):
            set_paragraph_text(para, "")


def _clone_row_after(table: Table, template_index: int, after_index: int) -> None:
    """克隆模板行并插到指定行之后。"""
    template_tr = deepcopy(table.rows[template_index]._tr)
    table.rows[after_index]._tr.addnext(template_tr)


def ensure_row_count(
    table: Table, target: RowsTarget, needed: int, spec: TableSpec, columns: list[tuple[str, str]]
) -> RowsTarget:
    """把数据区行数调整为 needed：不足则克隆模板行，多余则仅删除全空行。

    非空的多余行不删除（那是母本里已写好的内容），改为在返回值中保留并在渲染结果里提示溢出。
    """
    data_rows = list(target.data_rows)
    if needed > len(data_rows):
        for _ in range(needed - len(data_rows)):
            _clone_row_after(table, target.template_row, data_rows[-1])
            data_rows.append(data_rows[-1] + 1)
    elif needed < len(data_rows):
        for idx in reversed(data_rows[needed:]):
            if _row_text(table.rows[idx]):
                continue  # 母本该行已有内容，不删
            tr = table.rows[idx]._tr
            tr.getparent().remove(tr)
    return analyze_data_area(table, spec, columns)


def write_rows(
    table: Table,
    target: RowsTarget,
    records: list[dict[str, str]],
    columns: list[tuple[str, str]],
    spec: TableSpec,
) -> int:
    """按列映射写入数据行，返回实际写入行数。"""
    rows = list(target.data_rows)
    written = 0
    for offset, record in enumerate(records):
        if offset >= len(rows):
            break
        row_index = rows[offset]
        cells = row_cells(table, row_index)
        label_cell = cells[0] if cells else None
        for key, col_index in target.columns.items():
            if col_index >= len(cells):
                continue
            value = record.get(key)
            if value is None:
                continue
            set_cell_text(cells[col_index], value, style_from=label_cell)
        written += 1
    return written


def clear_rows(table: Table, target: RowsTarget, indexes: list[int]) -> None:
    """清空指定数据行（用于避免残留母本示例内容）。"""
    for idx in indexes:
        if idx in target.data_rows:
            _clear_row(table, idx)
