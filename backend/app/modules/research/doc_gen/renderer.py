"""把槽位值渲染进母本副本，产出初版 docx。

规则：
1. 任一锚点未命中 → 记 anchor_unresolved，绝不退化为「按关键词猜位置」。
2. 只写文字与克隆行，不改样式表、页眉页脚结构、列宽与合并方式。
3. 表格的序号、费用、占比、合计由 calc 计算，模型不参与算术。
"""

from __future__ import annotations

import io
import logging
import re
from collections.abc import Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.modules.research.doc_gen import calc, status
from app.modules.research.doc_gen.anchors import (
    CellTarget,
    GlobalTarget,
    ParagraphTarget,
    RowsTarget,
    analyze_data_area,
    resolve_anchor,
)
from app.modules.research.doc_gen.grid import cell_at_grid, normalize
from app.modules.research.doc_gen.template_spec import ColumnSpec, Slot, TemplateSpec
from app.modules.research.doc_gen.writer import (
    ensure_row_count,
    replace_keyword,
    row_cells,
    set_after_label,
    set_cell_text,
    write_multiline,
)

logger = logging.getLogger(__name__)

NumberExpects = ("number", "percent")

_TOC_TRAILING_PAGE = re.compile(r"\d+\s*$")
_TOC_HEADING = "目录"

# 正文样式的候选名（中英文母本）。
# 动态章节的正文段必须落到母本既有的正文样式上，不能沿用标题样式。
_BODY_STYLE_CANDIDATES = ("Normal", "正文", "Body Text", "Normal (Web)")


@dataclass(slots=True)
class SlotValue:
    """一个槽位的最终渲染值。"""

    text: str = ""
    rows: list[dict[str, str]] = field(default_factory=list)
    state: str = status.STATUS_OK


@dataclass(slots=True)
class SectionPlan:
    """一个动态章节实例的渲染计划。

    由流水线依据模板规格与落库的大纲组装；渲染器只负责「按顺序插标题、写内容」，
    不参与结构决策。
    """

    section_key: str
    fragment_key: str
    anchor_key: str  # 'ext:<扩展点 key>' 或父章节实例 key
    level: int
    title: str
    parent_section_key: str | None = None
    slots: list[Slot] = field(default_factory=list)
    on_empty: str = "keep_title"


@dataclass(slots=True)
class RenderReport:
    """渲染结果统计，供生成说明与接口返回使用。"""

    written: int = 0
    pending: list[str] = field(default_factory=list)
    drafts: list[str] = field(default_factory=list)
    unresolved: dict[str, str] = field(default_factory=dict)
    occupied: list[str] = field(default_factory=list)
    overflow: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    toc_paragraphs: int = 0
    sections_added: int = 0
    section_titles: list[str] = field(default_factory=list)
    sections_skipped: list[str] = field(default_factory=list)


def open_document(data: bytes) -> DocxDocument:
    """从字节流打开 docx/dotx。"""
    from docx import Document

    return Document(io.BytesIO(data))


def document_bytes(doc: DocxDocument) -> bytes:
    """序列化为 docx 字节。"""
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_document(
    master: bytes,
    spec: TemplateSpec,
    values: Mapping[str, SlotValue],
    sections: Sequence[SectionPlan] = (),
) -> tuple[bytes, RenderReport]:
    """渲染主入口：先填骨架槽位，再按大纲插入动态章节。"""
    doc = open_document(master)
    report = RenderReport()
    report.toc_paragraphs = apply_toc_notice(doc, spec.toc_notice)
    for slot in spec.slots:
        render_slot(doc, slot, values.get(slot.key) or SlotValue(text=status.pending()), report)
    if sections:
        _render_sections(doc, spec, sections, values, report)
    return document_bytes(doc), report


def render_slot(doc: DocxDocument, slot: Slot, value: SlotValue, report: RenderReport) -> None:
    """渲染单个槽位。"""
    if slot.kind == "table":
        _render_table_slot(doc, slot, value, report)
        return
    text = value.text.strip()
    if not text:
        text = status.pending()
    if slot.kind == "image":
        # 本版不解析/插入图片，只留文字占位，避免结构式与路线图被静默丢弃
        text = status.IMAGE_PLACEHOLDER.format(label=slot.label)
    if len(text) > slot.max_chars:
        report.overflow.append(f"{slot.key}（{len(text)} 字 > 上限 {slot.max_chars}）")
    if status.is_pending(text):
        report.pending.append(slot.key)
    elif value.state == status.STATUS_DRAFT:
        report.drafts.append(slot.key)
    resolved = 0
    for anchor in slot.anchors:
        try:
            target = resolve_anchor(doc, slot, anchor)
        except Exception as exc:  # noqa: BLE001 - 锚点未命中只记录，不中断整篇渲染
            report.unresolved[slot.key] = str(exc)
            continue
        if isinstance(target, CellTarget) and not slot.overwrite_cell and _cell_has_content(target):
            report.occupied.append(slot.key)
            continue
        _write_target(target, text)
        resolved += 1
    if resolved:
        report.written += 1


def _cell_has_content(target: CellTarget) -> bool:
    """目标单元格是否已有非占位内容。"""
    cell = cell_at_grid(target.table, target.row, target.col)
    if cell is None:
        return False
    text = normalize(cell.text)
    return bool(text) and not text.startswith(status.PENDING_PREFIX)


def _write_target(target: object, text: str) -> None:
    """把文本写入已定位的目标。"""
    if isinstance(target, CellTarget):
        cell = cell_at_grid(target.table, target.row, target.col)
        if cell is None:
            return
        cells = row_cells(target.table, target.row)
        set_cell_text(cell, text, style_from=cells[0] if cells else None)
        return
    if isinstance(target, ParagraphTarget):
        if target.mode == "after_label":
            set_after_label(target.paragraph, text)
        else:
            write_multiline(target.paragraph, text, mode="replace")
        return
    if isinstance(target, GlobalTarget):
        for para in target.paragraphs:
            replace_keyword(para, target.keyword, text)


def _render_table_slot(doc: DocxDocument, slot: Slot, value: SlotValue, report: RenderReport) -> None:
    """渲染表格槽位（含序号、计算列与合计行）。"""
    anchor = next(a for a in slot.anchors if a.type == "table_rows")
    try:
        target = resolve_anchor(doc, slot, anchor)
    except Exception as exc:  # noqa: BLE001
        report.unresolved[slot.key] = str(exc)
        return
    if not isinstance(target, RowsTarget):  # pragma: no cover - 防御
        report.unresolved[slot.key] = "定位配置异常，需复核模板"
        return
    spec = slot.table
    if spec is None:  # pragma: no cover - 模板校验已保证
        return
    records = [dict(row) for row in value.rows]
    if not records:
        report.pending.append(slot.key)
        return
    if len(records) > spec.row_limit:
        records = records[: spec.row_limit]
        report.overflow.append(f"{slot.key}（行数超过模板原件可承载上限 {spec.row_limit}）")
    columns = [(c.key, c.label) for c in slot.columns]
    _apply_computations(slot, records)
    try:
        target = analyze_data_area(target.table, spec, columns)
        _clear_mapped_cells(target)  # 先清空（含母本预填序号），空出的预留行才允许被裁掉
        target = ensure_row_count(target.table, target, len(records), spec, columns)
    except LookupError as exc:
        report.unresolved[slot.key] = f"数据区结构无法解析：{exc}"
        return
    for offset, record in enumerate(records):
        if offset >= len(target.data_rows):
            report.overflow.append(f"{slot.key}（记录数 {len(records)} 超过可用数据行 {len(target.data_rows)}）")
            break
        _write_record_row(target, offset, record)
    _write_total_row(slot, target, records)
    report.written += 1


def _apply_computations(slot: Slot, records: list[dict[str, str]]) -> None:
    """填写序号列与派生计算列。"""
    spec = slot.table
    columns = {c.key: c for c in slot.columns}
    sequence_key = spec.sequence_column if spec is not None else None
    for idx, record in enumerate(records):
        if sequence_key and sequence_key in columns:
            record[sequence_key] = str(idx + 1)
    # 先算乘积列，再算占比列：占比分母必须是全表合计，不能边算边累积
    for record in records:
        for key, column in columns.items():
            if column.compute == "product":
                product = calc.product([calc.to_number(record.get(op, "")) for op in column.operands])
                record[key] = calc.format_number(product) if product is not None else status.pending("用量或单价未提供")
    for record in records:
        for key, column in columns.items():
            if column.compute == "ratio":
                operand = column.operands[0] if column.operands else key
                part = calc.to_number(record.get(operand, ""))
                record[key] = calc.percent(part, _column_total(records, column.operands))


def _column_total(records: list[dict[str, str]], keys: list[str]) -> Decimal | None:
    """按操作数列求和（用于占比分母）。"""
    key = keys[0] if keys else ""
    return calc.sum_values([calc.to_number(record.get(key, "")) for record in records])


def _column_total_for(column: ColumnSpec, records: list[dict[str, str]]) -> Decimal | None:
    """合计行中某一列的总和。"""
    if column.compute == "ratio":
        return calc.sum_values([Decimal(100)] if records else [])
    return _column_total(records, [column.key])


def _clear_mapped_cells(target: RowsTarget) -> None:
    """清空数据区中受管列，避免母本预留内容（如预填序号）残留。"""
    for row_index in target.data_rows:
        for col_index in target.columns.values():
            cell = cell_at_grid(target.table, row_index, col_index)
            if cell is not None:
                set_cell_text(cell, "")


def _write_record_row(target: RowsTarget, offset: int, record: dict[str, str]) -> None:
    """写一条数据记录。"""
    row_index = target.data_rows[offset]
    cells = row_cells(target.table, row_index)
    label_cell = cells[0] if cells else None
    for key, col_index in target.columns.items():
        value = record.get(key)
        if value is None:
            continue
        cell = cell_at_grid(target.table, row_index, col_index)
        if cell is not None:
            set_cell_text(cell, value, style_from=label_cell)


def _write_total_row(slot: Slot, target: RowsTarget, records: list[dict[str, str]]) -> None:
    """合计行：数值列求和，占比列写 100%。"""
    if target.total_row is None or not records:
        return
    cells = row_cells(target.table, target.total_row)
    label_cell = cells[0] if cells else None
    for column in slot.columns:
        col_index = target.columns.get(column.key)
        if col_index is None or column.expects not in NumberExpects:
            continue
        if column.compute == "none":
            continue  # 合计行只写派生列（金额/占比），原始数值列求和没有业务含义
        cell = cell_at_grid(target.table, target.total_row, col_index)
        if cell is None:
            continue
        if column.compute == "ratio":
            set_cell_text(cell, "100%")
            continue
        total = _column_total_for(column, records)
        if total is not None:
            set_cell_text(cell, calc.format_number(total), style_from=label_cell)


def apply_toc_notice(doc: DocxDocument, notice: str) -> int:
    """母本目录为纯文本（无 TOC 域），页码无法自动刷新：去掉旧页码并插入提示。"""
    paragraphs = doc.paragraphs
    start = next((i for i, p in enumerate(paragraphs) if p.text.strip() == _TOC_HEADING), None)
    if start is None:
        return 0
    changed = 0
    last_entry = start
    for idx in range(start + 1, min(start + 60, len(paragraphs))):
        para = paragraphs[idx]
        text = para.text.strip()
        if not text:
            continue
        if not re.match(r"^\d+(\.\d+)*", text):
            break
        stripped = _TOC_TRAILING_PAGE.sub("", text).strip()
        if stripped != text:
            write_multiline(para, stripped, mode="replace")
            changed += 1
        last_entry = idx
    if changed:
        anchor = paragraphs[last_entry]
        clone = deepcopy(anchor._p)
        anchor._p.addnext(clone)
        notice_para: Paragraph = Paragraph(clone, anchor._parent)
        write_multiline(notice_para, notice, mode="replace", style_from=anchor)
    return changed


# ---------------------------------------------------------------------------
# 动态章节渲染
#
# 结构决策在大纲阶段已经完成，这里只做两件事：按顺序插入标题、写入章节内容。
# 两条硬约束：
#   1. 标题必须套用母本既有的 Heading 样式（样式不存在即显式失败，不降级成手工加粗），
#      编号交给 Word 的多级列表跟随样式自动生成；
#   2. 章节局部槽位写在本章节标题之后（section_relative），游标随写入前移，
#      父章节正文 → 子章节 → 下一个同级章节，母本中的后置内容自动顺延。
# ---------------------------------------------------------------------------


def _is_section_relative(slot: Slot) -> bool:
    """槽位是否使用「本章节内相对定位」：这类槽位在母本中不存在对应标题。"""
    return any(anchor.type == "section_relative" for anchor in slot.anchors)


def _slot_has_substance(value: SlotValue | None, slot: Slot) -> bool:
    """章节内容是否为实质内容（供 on_empty=skip_section 判定）。"""
    if value is None:
        return False
    if slot.kind == "table":
        return bool(value.rows)
    text = (value.text or "").strip()
    return bool(text) and not status.is_pending(text)


def _find_heading_paragraph(doc: DocxDocument, heading: str | None) -> Paragraph | None:
    """按标题文字定位母本中的挂载段落，优先命中标题样式的段落（避开目录里的同名条目）。"""
    if not heading:
        return None
    needle = normalize(heading)
    fallback: Paragraph | None = None
    for para in doc.paragraphs:
        if normalize(para.text) != needle:
            continue
        style_name = (para.style.name if para.style is not None else "") or ""
        if "heading" in style_name.lower() or "标题" in style_name:
            return para
        if fallback is None:
            fallback = para
    return fallback


def _heading_style_name(doc: DocxDocument, level: int) -> str | None:
    """取该层级的标题样式名；样式不存在时返回 None，由调用方记录未定位。"""
    exact = {f"Heading {level}", f"heading {level}", f"标题 {level}", f"标题{level}"}
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        name = str(style.name or "").strip()
        if name in exact:
            return name
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        if (getattr(style, "style_id", "") or "") == f"Heading{level}":
            name = str(style.name or "").strip()
            return name or None
    return None


def _body_style_name(doc: DocxDocument) -> str | None:
    """取一个正文样式用于动态章节正文段；找不到时退回默认样式。"""
    names = {str(p.style.name) for p in doc.paragraphs if p.style is not None and p.style.name}
    for candidate in _BODY_STYLE_CANDIDATES:
        if candidate in names:
            return candidate
    for name in sorted(names):
        lowered = name.lower()
        if "heading" not in lowered and "标题" not in name:
            return name
    return None


def _insert_heading_after(ref: Paragraph, style_name: str | None, title: str) -> Paragraph | None:
    """在参考段落后插入一个标题段落；样式缺失即返回 None。"""
    if not style_name:
        return None
    try:
        paragraph = ref.insert_paragraph_before(title, style=style_name)
    except KeyError:
        return None
    ref._p.addnext(paragraph._p)
    return paragraph


def _write_relative(
    cursor: Paragraph, body_style: str | None, slot: Slot, value: SlotValue, report: RenderReport
) -> Paragraph:
    """把章节内相对定位的槽位写在游标之后，返回新的游标段落。"""
    text = (value.text or "").strip()
    if slot.kind == "image":
        text = text or status.IMAGE_PLACEHOLDER.format(label=slot.label)
    if not text:
        text = status.pending()
    if len(text) > slot.max_chars:
        report.overflow.append(f"{slot.key}（{len(text)} 字 > 上限 {slot.max_chars}）")
    if status.is_pending(text):
        report.pending.append(slot.key)
    elif value.state == status.STATUS_DRAFT:
        report.drafts.append(slot.key)
    try:
        if body_style:
            paragraph = cursor.insert_paragraph_before("", style=body_style)
        else:
            paragraph = cursor.insert_paragraph_before("")
    except KeyError:
        paragraph = cursor.insert_paragraph_before("")
    cursor._p.addnext(paragraph._p)
    written = write_multiline(paragraph, text, mode="replace")
    report.written += 1
    return written[-1] if written else paragraph


def _render_sections(
    doc: DocxDocument,
    spec: TemplateSpec,
    sections: Sequence[SectionPlan],
    values: Mapping[str, SlotValue],
    report: RenderReport,
) -> None:
    """按大纲顺序插入动态章节（DFS 前序：父章节先于子章节，同级按 order 排列）。"""
    cursors: dict[str, Paragraph] = {}
    failed: set[str] = set()
    heading_styles: dict[int, str | None] = {}
    parents = {section.section_key: section.parent_section_key for section in sections}
    body_style = _body_style_name(doc)

    def resolve_cursor(key: str) -> Paragraph | None:
        if key in failed:
            return None
        cached = cursors.get(key)
        if cached is not None:
            return cached
        if not key.startswith("ext:"):
            report.unresolved[key] = f"章节挂载点 {key} 不可用（挂载章节未成功渲染）"
            failed.add(key)
            return None
        point = spec.extension_point(key[4:])
        heading = point.heading if point is not None else None
        para = _find_heading_paragraph(doc, heading)
        if para is None:
            report.unresolved[key] = f"扩展点挂载标题未找到：{heading!r}"
            failed.add(key)
            return None
        cursors[key] = para
        return para

    for section in sections:
        anchor = resolve_cursor(section.anchor_key)
        if anchor is None:
            continue
        substantive = sum(1 for slot in section.slots if _slot_has_substance(values.get(slot.key), slot))
        if section.slots and substantive == 0 and section.on_empty == "skip_section":
            report.sections_skipped.append(section.section_key)
            cursors[section.section_key] = anchor
            continue
        level = max(1, section.level)
        if level not in heading_styles:
            heading_styles[level] = _heading_style_name(doc, level)
        heading = _insert_heading_after(anchor, heading_styles[level], section.title)
        if heading is None:
            report.unresolved[section.section_key] = f"母本缺少 {level} 级标题样式，无法插入章节「{section.title}」"
            cursors[section.section_key] = anchor
            continue
        report.sections_added += 1
        report.section_titles.append(section.title)

        cursor = heading
        for slot in section.slots:
            value = values.get(slot.key) or SlotValue(text=status.pending())
            if _is_section_relative(slot):
                cursor = _write_relative(cursor, body_style, slot, value, report)
            else:
                render_slot(doc, slot, value, report)

        # 游标上溯：本节及其子树结束后，所有祖先的插入位置都要顺延到此处，
        # 否则下一个同级章节会插到本节子树之前。
        node: str | None = section.section_key
        while node is not None:
            cursors[node] = cursor
            node = parents.get(node)
        cursors[section.anchor_key] = cursor


def table_of(doc: DocxDocument, header_contains: list[str]) -> Table | None:
    """工具函数：按表头取表格（供测试与排错使用）。"""
    from app.modules.research.doc_gen.anchors import find_table

    try:
        return find_table(doc, header_contains)
    except LookupError:
        return None


def slot_value_from_mapping(data: Mapping[str, Any]) -> SlotValue:
    """把服务层结果转成渲染值。"""
    if isinstance(data, SlotValue):
        return data
    if isinstance(data, list):
        rows = [{str(k): str(v) for k, v in dict(item).items()} for item in data]
        return SlotValue(rows=rows, state=status.STATUS_OK)
    return SlotValue(text=str(data or ""))
