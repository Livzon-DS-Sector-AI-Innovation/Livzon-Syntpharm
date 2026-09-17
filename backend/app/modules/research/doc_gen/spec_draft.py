"""从 Word 母本自动草拟槽位定义（上传时不要求用户选择任何配置）。

能自动推断的是**定位**（anchor）：表头+空数据行、空值标签段落、彩色提示语段落、
章节标题后的空白正文、页眉受控字段。母本里不存在的是**语义**，一律取保守默认：

- `source_roles` 只允许 material（不引用文献）；
- 表格列不做 `compute`（不猜合计/占比）；
- 字段与表格不允许 AI 草稿；只有「提示语段落 / 空章节正文」这类明显是「此处请撰写正文」
  的位置才允许草稿（前缀会被标成【AI草稿，需确认】）；
- `required` 一律 False，取不到就写 `[待补充：原因]`。

这样自动识别出来的配置可以直接用，且不违反「没有依据就不写」与「样式零改动」。
"""

from __future__ import annotations

import hashlib
import logging
from typing import Any

from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

from app.modules.research.doc_gen.anchors import header_paragraphs, run_color
from app.modules.research.doc_gen.grid import cell_text, distinct_row_cells, normalize
from app.modules.research.doc_gen.renderer import open_document
from app.modules.research.doc_gen.template_spec import (
    Anchor,
    ColumnSpec,
    Slot,
    TableSpec,
    TemplateSpec,
)

logger = logging.getLogger(__name__)

DEFAULT_HINT_COLORS: tuple[str, ...] = ("00B0F0",)
# 页眉受控字段：key 必须与 job.meta 的键一致，否则渲染阶段取不到值
_HEADER_FIELDS: tuple[tuple[str, str, str], ...] = (
    ("doc_code", "受控编码", "编码"),
    ("doc_version", "版本号", "版本号"),
)
_TOTAL_LABELS = ("总价", "合计", "总计")
_SEQ_HEADERS = ("序号",)
_LABEL_TAILS = ("：", ":")
_IMAGE_LABELS = ("结构式", "图片", "图谱")
_HEADING_HINTS = ("heading", "标题")
_TOC_HINTS = ("toc",)
_MAX_SLOTS = 120


def _is_heading(para: Paragraph) -> bool:
    """按样式名判断标题（兼容中英文 Word 样式）。"""
    name = str(getattr(para.style, "name", "") or "").lower()
    return any(hint in name for hint in _HEADING_HINTS)


def _is_toc(para: Paragraph) -> bool:
    """目录条目（样式名含 toc），既不是标题也不该当正文。"""
    name = str(getattr(para.style, "name", "") or "").lower()
    return any(hint in name for hint in _TOC_HINTS)


def _short(text: str, limit: int = 100) -> str:
    """压缩空白并截断，用于 label / key 生成。"""
    joined = " ".join(text.split())
    return joined[:limit]


def _spec_code(data: bytes) -> str:
    """按母本内容生成稳定 code：同一份文件重复上传会命中同一套定义。"""
    return f"auto_{hashlib.sha1(data).hexdigest()[:10]}"


def _scan_header_fields(doc: DocxDocument) -> list[Slot]:
    """页眉里的受控字段（编码/版本号），值来自表单而非资料。"""
    slots: list[Slot] = []
    for key, label, needle in _HEADER_FIELDS:
        if not header_paragraphs(doc, needle):
            continue
        slots.append(
            Slot(
                key=key,
                label=label,
                from_meta=True,
                max_chars=40,
                anchors=[Anchor(type="header_field", header_contains=needle)],
            )
        )
    return slots


def _scan_tables(doc: DocxDocument) -> list[Slot]:
    """识别「表头 + 空数据行」的可填表格。

    只接受**表头单元格全部非空**的表（合并/多行表头无法安全推断列语义，宁可漏也不猜），
    且必须存在一行「除序号列外全空」，以保证这是待填表而不是固定结构表。
    """
    slots: list[Slot] = []
    index = 0
    for table in doc.tables:
        rows = list(table.rows)
        if len(rows) < 2:
            continue
        headers = [cell_text(c) for c in distinct_row_cells(rows[0])]
        if len(headers) < 2 or any(not h for h in headers):
            continue
        columns: list[ColumnSpec] = []
        seq_key: str | None = None
        for pos, text in enumerate(headers):
            key = f"c{pos + 1}"
            columns.append(ColumnSpec(key=key, label=text, max_chars=200))
            if text in _SEQ_HEADERS:
                seq_key = key
        total_label: str | None = None
        data_rows: list[list[str]] = []
        for row in rows[1:]:
            texts = [cell_text(c) for c in distinct_row_cells(row)]
            first = texts[0] if texts else ""
            if first and any(label in first for label in _TOTAL_LABELS):
                total_label = first
                continue
            data_rows.append(texts)
        fillable = any(
            all(not text for pos, text in enumerate(texts) if f"c{pos + 1}" != seq_key) for texts in data_rows
        )
        if not fillable:
            continue
        index += 1
        slots.append(
            Slot(
                key=f"table_{index}",
                label=_short(" / ".join(headers), 150),
                kind="table",
                table=TableSpec(
                    header_rows=1,
                    row_limit=60,
                    sequence_column=seq_key,
                    total_row_label=total_label,
                ),
                columns=columns,
                anchors=[Anchor(type="table_rows", table_header=headers)],
            )
        )
    return slots


def _scan_label_paragraphs(doc: DocxDocument) -> list[Slot]:
    """正文里「标签：」且冒号后为空的位置（如「制剂/规格：」）。"""
    slots: list[Slot] = []
    index = 0
    for para in doc.paragraphs:
        text = normalize(para.text)
        if not text or len(text) > 60 or not text.endswith(_LABEL_TAILS):
            continue
        if _is_heading(para) or _is_toc(para):
            continue
        index += 1
        raw = para.text.strip()
        if any(hint in raw for hint in _IMAGE_LABELS):
            # 结构式/图片这类位置只留文字占位，人工贴图
            slots.append(
                Slot(
                    key=f"image_{index}",
                    label=_short(raw),
                    kind="image",
                    anchors=[Anchor(type="image_placeholder", paragraph_label=raw)],
                )
            )
            continue
        slots.append(
            Slot(
                key=f"label_{index}",
                label=_short(raw),
                kind="field",
                max_chars=300,
                anchors=[Anchor(type="paragraph_after_label", paragraph_label=raw, paragraph_match="exact")],
            )
        )
    return slots


def _scan_hint_paragraphs(doc: DocxDocument, colors: tuple[str, ...]) -> list[Slot]:
    """彩色提示语段落（母本用蓝字告诉人「这里要写什么」）→ 整段替换。"""
    wanted = {c.upper() for c in colors}
    slots: list[Slot] = []
    index = 0
    for para in doc.paragraphs:
        text = normalize(para.text)
        if not text or _is_heading(para) or _is_toc(para):
            continue
        if not any(run_color(run).upper() in wanted for run in para.runs):
            continue
        index += 1
        raw = para.text.strip()
        slots.append(
            Slot(
                key=f"paragraph_{index}",
                label=_short(raw),
                kind="paragraph",
                draft_allowed=True,
                max_chars=2000,
                anchors=[Anchor(type="replace_paragraph", paragraph_label=raw, paragraph_match="contains")],
            )
        )
    return slots


def _scan_section_bodies(doc: DocxDocument) -> list[Slot]:
    """章节标题下没有预留正文的位置 → 在标题后补写正文。"""
    paras = list(doc.paragraphs)
    slots: list[Slot] = []
    index = 0
    for pos, para in enumerate(paras):
        if not _is_heading(para):
            continue
        raw = para.text.strip()
        if not raw or normalize(raw) == "目录":
            continue
        following = next((p for p in paras[pos + 1 :] if p.text.strip()), None)
        if following is None or not _is_heading(following):
            continue  # 标题下已有正文，交给其他规则处理
        index += 1
        slots.append(
            Slot(
                key=f"section_{index}",
                label=_short(raw),
                kind="paragraph",
                draft_allowed=True,
                max_chars=2000,
                anchors=[Anchor(type="section_body", heading=raw)],
            )
        )
    return slots


def draft_spec_from_bytes(
    data: bytes,
    *,
    name: str = "",
    code: str | None = None,
    version: str = "01",
    stage: str = "",
    hint_colors: tuple[str, ...] | None = None,
) -> TemplateSpec:
    """扫描一份 Word 母本，产出可直接使用的槽位定义。

    打开失败会抛出异常（调用方需捕获并给出「文件损坏」类提示），
    其余情况即使一个槽位都没识别到也返回合法但为空的规格——由调用方决定是否拒绝。
    """
    doc = open_document(data)
    colors = tuple(hint_colors or DEFAULT_HINT_COLORS)
    slots: list[Slot] = []
    slots += _scan_header_fields(doc)
    slots += _scan_tables(doc)
    slots += _scan_label_paragraphs(doc)
    slots += _scan_hint_paragraphs(doc, colors)
    slots += _scan_section_bodies(doc)

    deduped: list[Slot] = []
    seen: set[tuple[str, str]] = set()
    for slot in slots:
        marker = (slot.kind, slot.label)
        if marker in seen:
            continue
        seen.add(marker)
        deduped.append(slot)
        if len(deduped) >= _MAX_SLOTS:
            break

    spec = TemplateSpec(
        code=code or _spec_code(data),
        name=name or "自动识别的文档模板",
        version=version,
        stage=stage,
        master_asset="",
        description="由上传的 Word 母本自动识别填充项生成",
        hint_colors=list(colors),
        unfilled_notes=[
            "填充项由母本结构自动识别，语义为保守默认：不引用文献、不做表格计算、"
            "取不到依据一律标 [待补充]",
        ],
        slots=deduped,
    )
    logger.info(
        "母本槽位自动识别完成",
        extra={"code": spec.code, "slots": len(deduped), "module_code": "research"},
    )
    return spec


def summarize_spec(spec: TemplateSpec) -> dict[str, Any]:
    """概况（用于日志与接口回显）。"""
    kinds: dict[str, int] = {}
    for slot in spec.slots:
        kinds[slot.kind] = kinds.get(slot.kind, 0) + 1
    return {
        "code": spec.code,
        "name": spec.name,
        "version": spec.version,
        "slot_count": len(spec.slots),
        "kinds": kinds,
    }


__all__ = ["DEFAULT_HINT_COLORS", "draft_spec_from_bytes", "summarize_spec"]
