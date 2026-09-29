"""docx 母本 → 全内容 Markdown：模板 Markdown 弹窗数据源（标题/段落/列表/表格按文档流）。"""

from __future__ import annotations

import io
from collections.abc import Callable
from types import SimpleNamespace
from typing import Any
from uuid import UUID, uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE

from app.core.exceptions import BadRequestException, NotFoundException
from app.modules.research.doc_gen import service, spec_source
from app.modules.research.doc_gen.docx_markdown import render_docx_markdown
from app.modules.research.doc_gen.template_spec import Anchor, Slot, TemplateSpec


def _docx_bytes(build: Callable[[DocxDocument], None]) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_full(doc: DocxDocument) -> None:
    doc.add_heading("第一章 概述", level=1)
    doc.add_paragraph("这是正文段落，描述研究背景。")
    doc.add_heading("1.1 目的", level=2)
    doc.add_paragraph("列表项样式段落", style="List Bullet")
    table = doc.add_table(rows=2, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "规格|单位"
    hdr[2].text = "结果"
    body = table.rows[1].cells
    body[0].text = "收率"
    body[1].text = "98.5%\n（HPLC）"
    body[2].text = "合格"


def test_render_headings_paragraphs_and_tables() -> None:
    """标题按样式映射 #，段落原样，表格转 GFM 且首行后插分隔行。"""
    markdown = render_docx_markdown(_docx_bytes(_build_full))
    assert "# 第一章 概述" in markdown
    assert "## 1.1 目的" in markdown
    assert "这是正文段落，描述研究背景。" in markdown
    assert "- 列表项样式段落" in markdown
    assert "| 项目 | 规格\\|单位 | 结果 |" in markdown
    assert "| --- | --- | --- |" in markdown
    assert "98.5%<br>（HPLC）" in markdown
    # 文档流顺序：标题在段落前、段落在表格前
    assert markdown.index("# 第一章 概述") < markdown.index("这是正文段落") < markdown.index("| 项目 |")


def test_render_empty_document_returns_blank() -> None:
    markdown = render_docx_markdown(_docx_bytes(lambda doc: None))
    assert markdown == ""


def test_render_chinese_heading_style() -> None:
    """中文版式样式名「标题 1」同样映射为一级标题。"""
    style = Document().styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)

    def build(doc: DocxDocument) -> None:
        doc.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("中文标题", style="标题 1")
        assert style.name == "标题 1"

    markdown = render_docx_markdown(_docx_bytes(build))
    assert "# 中文标题" in markdown


def _build_literal_markers(doc: DocxDocument) -> None:
    doc.add_paragraph("质量收率＝****中间体1产出量÷****投入量*100%")
    doc.add_paragraph("供应商 <b>公司</b> 与 [待定] 及 _下划线_ 混排")
    table = doc.add_table(rows=2, cols=1)
    table.rows[0].cells[0].text = "涂黑"
    table.rows[1].cells[0].text = "****匿名"


def test_render_escapes_literal_markdown_in_text() -> None:
    """母本正文/单元格里的字面星号、尖括号、下划线、方括号被转义，不会被误当强调或 HTML 语法。"""
    markdown = render_docx_markdown(_docx_bytes(_build_literal_markers))
    # 成对/连续星号转义为 \*，星号不再被吃掉、也不会错加粗斜体
    assert r"\*\*\*\*中间体1产出量" in markdown
    assert r"\*100%" in markdown
    # 尖括号转义，杜绝 rehype-raw 下把正文当 HTML 注入
    assert r"\<b\>公司\</b\>" in markdown
    assert r"\[待定\]" in markdown
    assert r"\_下划线\_" in markdown
    # 单元格同样转义：不应残留连续的裸星号
    assert r"\*\*\*\*匿名" in markdown
    assert "****匿名" not in markdown


# ===== service.template_markdown：母本优先、骨架回退 =====


def _fake_template() -> SimpleNamespace:
    return SimpleNamespace(id=uuid4(), name="测试模板", template_code="tc", file_object_key="key.docx")


def _fake_spec() -> TemplateSpec:
    return TemplateSpec(
        code="test_md",
        name="骨架模板",
        version="v1",
        master_asset="tech_research_report.dotx",
        slots=[
            Slot(
                key="project_name",
                label="项目名称",
                anchors=[Anchor(type="paragraph_after_label", paragraph_label="项目名称")],
            )
        ],
        fragments=[],
    )


def _patch(
    monkeypatch: pytest.MonkeyPatch, template: Any, spec: TemplateSpec | None, master: bytes | Exception
) -> None:
    async def _load(_session: Any, _tid: UUID) -> Any:
        return template

    async def _resolve(_session: Any, _template: Any) -> TemplateSpec:
        if spec is None:
            raise KeyError("no spec")
        return spec

    async def _master(_template: Any, _spec: Any) -> bytes:
        if isinstance(master, Exception):
            raise master
        return master

    monkeypatch.setattr(service, "_load_template_or_404", _load)
    monkeypatch.setattr(spec_source, "resolve_for_template", _resolve)
    monkeypatch.setattr(service, "_template_master_bytes", _master)


async def test_service_prefers_docx_full_content(monkeypatch: pytest.MonkeyPatch) -> None:
    data = _docx_bytes(_build_full)
    _patch(monkeypatch, _fake_template(), _fake_spec(), data)
    result = await service.template_markdown(None, uuid4())  # type: ignore[arg-type]
    assert result["source"] == "docx"
    assert "# 第一章 概述" in result["markdown"]
    # 骨架元信息不在全内容里；slot_count 仍按 spec 回填供标签展示
    assert "<!--slot:" not in result["markdown"]
    assert result["slot_count"] == 1


async def test_service_falls_back_to_spec_without_master(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch(monkeypatch, _fake_template(), _fake_spec(), NotFoundException("无母本"))
    result = await service.template_markdown(None, uuid4())  # type: ignore[arg-type]
    assert result["source"] == "spec"
    assert "<!--slot:project_name-->" in result["markdown"]


async def test_service_falls_back_when_docx_broken(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch(monkeypatch, _fake_template(), _fake_spec(), b"not a docx at all")
    result = await service.template_markdown(None, uuid4())  # type: ignore[arg-type]
    assert result["source"] == "spec"


async def test_service_raises_when_nothing_available(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch(monkeypatch, _fake_template(), None, NotFoundException("无母本"))
    with pytest.raises(BadRequestException):
        await service.template_markdown(None, uuid4())  # type: ignore[arg-type]


async def test_service_returns_slot_semantics(monkeypatch: pytest.MonkeyPatch) -> None:
    """语义清单随 markdown 一并返回：AI 增强/人工维护的成果在「填写项语义」视图可核对。"""
    spec = _fake_spec()
    spec.slots[0].search_terms = ["课题名称", "项目代号"]
    spec.slots[0].query_hint = "立项书封面全称"
    spec.slots[0].review_state = "needs_review"
    _patch(monkeypatch, _fake_template(), spec, NotFoundException("无母本"))
    result = await service.template_markdown(None, uuid4())  # type: ignore[arg-type]
    assert result["slots"][0]["key"] == "project_name"
    assert result["slots"][0]["search_terms"] == ["课题名称", "项目代号"]
    assert result["slots"][0]["query_hint"] == "立项书封面全称"
    assert result["slots"][0]["review_state"] == "needs_review"
    assert result["needs_review"] == 1
