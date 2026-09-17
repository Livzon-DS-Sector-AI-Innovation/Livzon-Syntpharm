"""动态章节测试：触发判定、实例化展开、标题合成、多层级渲染。"""

from __future__ import annotations

import io
from typing import Any

import pytest
from docx import Document

from app.modules.research.doc_gen import renderer, status
from app.modules.research.doc_gen.extraction import SlotResult
from app.modules.research.doc_gen.parsing import TextBlock
from app.modules.research.doc_gen.sections import (
    build_trigger_context,
    compose_title,
    evaluate_trigger,
    instantiate,
    recompose_title,
)
from app.modules.research.doc_gen.template_spec import (
    Anchor,
    ExtensionPoint,
    SectionFragment,
    Slot,
    TemplateSpec,
    Trigger,
)


def _slot(key: str, kind: str = "field", **kwargs: Any) -> Slot:
    return Slot(key=key, label=key, kind=kind, anchors=[Anchor(type="section_relative")], **kwargs)


def _spec(**overrides: Any) -> TemplateSpec:
    base: dict[str, Any] = {
        "code": "t",
        "name": "测试模板",
        "version": "1",
        "master_asset": "x.docx",
        "slots": [],
        "extension_points": [
            ExtensionPoint(key="ep1", heading="工艺研究", level=1, fragments=["route_study"], max_instances=5)
        ],
        "fragments": [
            SectionFragment(
                key="route_study",
                level=1,
                title_pool=["路线{repeat:route_name}研究"],
                repeat_over="route_list",
                repeat_field="route_name",
                slots=[_slot("yield")],
            )
        ],
    }
    base.update(overrides)
    return TemplateSpec(**base)


def _ctx(results: dict[str, SlotResult] | None = None, text: str = "", meta: dict[str, Any] | None = None) -> Any:
    blocks = [TextBlock(file_id="m1", page=1, index=0, text=text, kind="paragraph")] if text else []
    return build_trigger_context(results=results or {}, blocks=blocks, roles={}, meta=meta or {})


# ── 模板规格校验 ─────────────────────────────────────────────


def test_title_pool_rejects_numbering_prefix() -> None:
    """标题池禁止含编号前缀：编号由 Word 样式生成，写进文字会双编号。"""
    with pytest.raises(Exception, match="编号前缀"):
        SectionFragment(key="x", title_pool=["1. 路线研究"], slots=[])
    with pytest.raises(Exception, match="编号前缀"):
        SectionFragment(key="x", title_pool=["三、杂质研究"], slots=[])


def test_title_pool_rejects_unknown_variable() -> None:
    with pytest.raises(Exception, match="不在白名单"):
        SectionFragment(key="x", title_pool=["{ai:xxx}研究"], slots=[])


def test_extension_point_requires_exactly_one_parent() -> None:
    with pytest.raises(Exception, match="只能声明 heading 或 fragment 之一"):
        ExtensionPoint(key="ep", level=1, fragments=["a"])
    with pytest.raises(Exception, match="只能声明 heading 或 fragment 之一"):
        ExtensionPoint(key="ep", heading="甲", fragment="a", level=1, fragments=["a"])


def test_fragment_cycle_is_rejected() -> None:
    """片段循环引用必须被拦掉。

    层级约束（子扩展点 level = 父 level + 1）使模板级循环在结构上不可达，
    因此这里直接测环检测器本身，作为将来放宽层级规则时的兜底。
    """
    with pytest.raises(ValueError, match="循环引用"):
        TemplateSpec._assert_acyclic({"a": ["b"], "b": ["a"]})
    TemplateSpec._assert_acyclic({"a": ["b"], "b": []})  # 无环不抛


def test_extension_point_level_must_match_fragment() -> None:
    with pytest.raises(Exception, match="层级"):
        TemplateSpec(
            code="t",
            name="t",
            version="1",
            master_asset="x",
            slots=[],
            extension_points=[ExtensionPoint(key="ep", heading="甲", level=1, fragments=["a"])],
            fragments=[SectionFragment(key="a", level=2, title_pool=["甲"], slots=[])],
        )


# ── 触发判定 ─────────────────────────────────────────────


def test_trigger_slot_filled() -> None:
    ctx = _ctx(results={"a": SlotResult(key="a", text="有值", state=status.STATUS_OK)})
    hit, _ = evaluate_trigger(Trigger(op="slot_filled", key="a"), ctx)
    assert hit is True
    ctx2 = _ctx(results={"a": SlotResult(key="a", text=status.pending(), state=status.STATUS_PENDING)})
    hit2, _ = evaluate_trigger(Trigger(op="slot_filled", key="a"), ctx2)
    assert hit2 is False


def test_trigger_table_rows_and_keyword() -> None:
    ctx = _ctx(results={"t": SlotResult(key="t", rows=[{"a": "1"}, {"a": "2"}])}, text="本品含有关物质杂质")
    hit, _ = evaluate_trigger(Trigger(op="table_rows_at_least", key="t", count=2), ctx)
    assert hit is True
    hit2, _ = evaluate_trigger(Trigger(op="keyword_found", keywords=["杂质"]), ctx)
    assert hit2 is True
    hit3, _ = evaluate_trigger(Trigger(op="keyword_found", keywords=["基因毒"]), ctx)
    assert hit3 is False


def test_trigger_combinators() -> None:
    ctx = _ctx(results={"a": SlotResult(key="a", text="v", state=status.STATUS_OK)})
    all_of = Trigger(op="all_of", of=[Trigger(op="slot_filled", key="a"), Trigger(op="slot_filled", key="b")])
    hit, trace = evaluate_trigger(all_of, ctx)
    assert hit is False and len(trace["children"]) == 2
    any_of = Trigger(op="any_of", of=[Trigger(op="slot_filled", key="a"), Trigger(op="slot_filled", key="b")])
    assert evaluate_trigger(any_of, ctx)[0] is True
    neg = Trigger(op="not", of=[Trigger(op="slot_filled", key="b")])
    assert evaluate_trigger(neg, ctx)[0] is True


def test_trigger_error_degrades_to_not_fired() -> None:
    """求值异常一律判为未触发，绝不让单个片段拖垮任务。"""
    ctx = _ctx()
    hit, trace = evaluate_trigger(Trigger(op="table_rows_at_least", key="missing"), ctx)
    assert hit is False


# ── 实例化与展开 ─────────────────────────────────────────────


def test_instantiate_repeats_per_row_and_expands_slots() -> None:
    spec = _spec()
    ctx = _ctx(results={"route_list": SlotResult(key="route_list", rows=[{"route_name": "A"}, {"route_name": "B"}])})
    outcome = instantiate(spec, ctx)
    assert [i.section_key for i in outcome.instances] == ["route_study__0", "route_study__1"]
    assert [i.title for i in outcome.instances] == ["路线A研究", "路线B研究"]
    assert outcome.instances[0].slot_keys == ["route_study__0.yield"]
    assert outcome.instances[1].slot_keys == ["route_study__1.yield"]


def test_instantiate_skips_when_trigger_not_hit() -> None:
    spec = _spec(
        fragments=[
            SectionFragment(
                key="impurity",
                level=1,
                title_pool=["杂质研究"],
                trigger=Trigger(op="keyword_found", keywords=["基因毒"]),
                slots=[_slot("note", kind="paragraph")],
            )
        ],
        extension_points=[ExtensionPoint(key="ep1", heading="工艺研究", level=1, fragments=["impurity"])],
    )
    outcome = instantiate(spec, _ctx(text="本品为普通工艺"))
    assert outcome.instances == []


def test_instantiate_caps_instances() -> None:
    spec = _spec()
    spec.extension_points[0].max_instances = 2
    rows = [{"route_name": chr(65 + i)} for i in range(5)]
    outcome = instantiate(spec, _ctx(results={"route_list": SlotResult(key="route_list", rows=rows)}))
    assert len(outcome.instances) == 2
    assert any("上限" in w or "截断" in w for w in outcome.warnings)


def test_instantiate_multilevel_unique_keys() -> None:
    """多层级：同一片段挂在不同父章节下，section_key 必须全局唯一。"""
    spec = _spec(
        fragments=[
            SectionFragment(
                key="route_study",
                level=1,
                title_pool=["路线{repeat:route_name}研究"],
                repeat_over="route_list",
                repeat_field="route_name",
                slots=[_slot("yield")],
                extension_points=[
                    ExtensionPoint(key="ep_sub", fragment="route_study", level=2, fragments=["impurity"])
                ],
            ),
            SectionFragment(
                key="impurity",
                level=2,
                title_pool=["杂质谱"],
                trigger=Trigger(op="keyword_found", keywords=["杂质"]),
                slots=[_slot("note", kind="paragraph")],
            ),
        ]
    )
    ctx = _ctx(
        results={"route_list": SlotResult(key="route_list", rows=[{"route_name": "A"}, {"route_name": "B"}])},
        text="含杂质",
    )
    outcome = instantiate(spec, ctx)
    keys = [i.section_key for i in outcome.instances]
    assert len(keys) == len(set(keys)), f"section_key 冲突：{keys}"
    # DFS 前序：父章节紧跟其子章节
    assert keys == ["route_study__0", "impurity__0", "route_study__1", "impurity__1"]
    children = [i for i in outcome.instances if i.fragment_key == "impurity"]
    assert {c.parent_section_key for c in children} == {"route_study__0", "route_study__1"}


# ── 标题合成 ─────────────────────────────────────────────


def test_compose_title_variables() -> None:
    fragment = SectionFragment(key="f", title_pool=["路线{repeat:route_name}研究"], repeat_field="route_name", slots=[])
    title, warnings = compose_title(fragment, 0, {"route_name": "A"}, _ctx())
    assert title == "路线A研究" and warnings == []


def test_recompose_title_requires_pool_member() -> None:
    fragment = SectionFragment(key="f", title_pool=["甲", "乙"], slots=[])
    with pytest.raises(ValueError, match="预设池"):
        recompose_title(fragment, "f__0", "池外标题", results={}, meta={})
    title, _ = recompose_title(fragment, "f__0", "乙", results={}, meta={})
    assert title == "乙"


# ── 渲染 ─────────────────────────────────────────────


def _master() -> bytes:
    doc = Document()
    doc.add_heading("工艺研究", level=1)
    doc.add_paragraph("正文占位")
    doc.add_heading("结论", level=1)
    doc.add_paragraph("结论占位")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plans(outcome: Any, spec: TemplateSpec) -> list[renderer.SectionPlan]:
    plans: list[renderer.SectionPlan] = []
    for i in outcome.instances:
        fragment = spec.fragment(i.fragment_key)
        plans.append(
            renderer.SectionPlan(
                section_key=i.section_key,
                fragment_key=i.fragment_key,
                anchor_key=i.anchor_key,
                level=i.level,
                title=i.title,
                parent_section_key=i.parent_section_key,
                slots=i.slots,
                on_empty=fragment.on_empty if fragment is not None else "keep_title",
            )
        )
    return plans


def test_render_inserts_headings_with_style() -> None:
    spec = _spec()
    ctx = _ctx(results={"route_list": SlotResult(key="route_list", rows=[{"route_name": "A"}, {"route_name": "B"}])})
    outcome = instantiate(spec, ctx)
    values = {
        "route_study__0.yield": renderer.SlotValue(text="85%"),
        "route_study__1.yield": renderer.SlotValue(text="72%"),
    }
    data, report = renderer.render_document(_master(), spec, values, _plans(outcome, spec))
    doc = renderer.open_document(data)
    by_text = {p.text.strip(): p for p in doc.paragraphs}
    assert "路线A研究" in by_text and "路线B研究" in by_text
    # 标题必须套用母本标题样式（编号由样式多级列表生成）
    assert "heading" in (by_text["路线A研究"].style.name or "").lower()
    assert report.sections_added == 2
    # 内容写在本章节标题之后
    texts = [p.text for p in doc.paragraphs]
    assert texts.index("85%") > texts.index("路线A研究")
    assert texts.index("路线B研究") > texts.index("85%")


def test_render_skip_section_when_empty() -> None:
    spec = _spec()
    spec.fragments[0].on_empty = "skip_section"
    ctx = _ctx(results={"route_list": SlotResult(key="route_list", rows=[{"route_name": "A"}])})
    outcome = instantiate(spec, ctx)
    data, report = renderer.render_document(_master(), spec, {}, _plans(outcome, spec))
    doc = renderer.open_document(data)
    assert "路线A研究" not in {p.text.strip() for p in doc.paragraphs}
    assert report.sections_skipped == ["route_study__0"]


def test_render_missing_anchor_heading_fails_visibly() -> None:
    """扩展点挂载标题在母本中不存在时，必须显式记入未定位，而不是静默丢章节。"""
    spec = _spec()
    spec.extension_points[0].heading = "不存在的章节"
    ctx = _ctx(results={"route_list": SlotResult(key="route_list", rows=[{"route_name": "A"}])})
    outcome = instantiate(spec, ctx)
    _data, report = renderer.render_document(_master(), spec, {}, _plans(outcome, spec))
    assert "ext:ep1" in report.unresolved
