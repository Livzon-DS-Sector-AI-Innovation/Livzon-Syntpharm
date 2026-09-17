"""doc_gen 资料解析（parsing.py）单元测试。OCR 全部 monkeypatch 打桩，禁止真实初始化。"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from app.modules.research.doc_gen.parsing import (
    ParseResult,
    parse_file,
    supported_extensions,
    supported_formats,
)


class FakeOCR:
    """假 OCR 服务：返回固定文本，记录调用。"""

    def __init__(self, text: str = "结晶温度 60\n收率 92%") -> None:
        self.text = text
        self.calls: list[Any] = []

    def extract_text(self, image_input: Any) -> str:
        self.calls.append(image_input)
        return self.text


def _make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("第一章 工艺概述", level=1)
    doc.add_paragraph("本文件描述阿莫西林结晶工艺。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "结晶温度"
    table.cell(1, 1).text = "60℃"
    doc.save(str(path))


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "工艺"
    ws.append(["参数", "数值"])
    ws.append(["结晶温度", 60])
    ws.append([None, None])
    ws2 = wb.create_sheet("设备")
    ws2.append(["设备名"])
    wb.save(str(path))


async def test_parse_docx(tmp_path: Path) -> None:
    path = tmp_path / "spec.docx"
    _make_docx(path)
    result = await parse_file(path, file_id="abc123")

    assert isinstance(result, ParseResult)
    assert result.warnings == []
    assert result.ocr_used is False
    assert all(b.file_id == "abc123" for b in result.blocks)
    assert all(b.page == 1 for b in result.blocks)  # docx 无页概念统一填 1
    assert [b.index for b in result.blocks] == list(range(len(result.blocks)))  # index 单调递增
    assert result.char_count == sum(len(b.text) for b in result.blocks)

    kinds = [(b.kind, b.text) for b in result.blocks]
    assert ("heading", "第一章 工艺概述") in kinds
    assert ("paragraph", "本文件描述阿莫西林结晶工艺。") in kinds
    assert ("table_row", "参数 | 数值") in kinds
    assert ("table_row", "结晶温度 | 60℃") in kinds


async def test_parse_docx_dotx_extension(tmp_path: Path) -> None:
    path = tmp_path / "template.dotx"
    _make_docx(path)
    result = await parse_file(path, file_id="t1")
    assert any(b.kind == "heading" for b in result.blocks)


async def test_parse_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    _make_xlsx(path)
    result = await parse_file(path, file_id="x1")

    assert result.warnings == []
    assert all(b.kind == "sheet_row" for b in result.blocks)
    # 空行被跳过：工艺 sheet 2 行 + 设备 sheet 1 行
    assert [b.text for b in result.blocks] == ["参数 | 数值", "结晶温度 | 60", "设备名"]
    # page 为 sheet 序号，从 1 递增
    assert [b.page for b in result.blocks] == [1, 1, 2]
    assert result.page_count == 2


async def test_parse_txt(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"
    path.write_text("第一段 内容\n继续第一段\n\n第二段 内容\n", encoding="utf-8")
    result = await parse_file(path, file_id="txt1")

    assert [b.text for b in result.blocks] == ["第一段 内容 继续第一段", "第二段 内容"]
    assert all(b.kind == "paragraph" and b.page == 1 for b in result.blocks)


async def test_unknown_extension_is_safe(tmp_path: Path) -> None:
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK\x03\x04junk")
    result = await parse_file(path, file_id="z1")

    assert result.blocks == []
    assert result.page_count == 0 and result.char_count == 0
    assert any("不支持的文件类型" in w for w in result.warnings)


async def test_missing_pdf_does_not_raise(tmp_path: Path) -> None:
    result = await parse_file(tmp_path / "not_exist.pdf", file_id="p1")
    assert result.blocks == []
    assert result.aborted == "not_found"
    assert any("文件不存在" in w for w in result.warnings)


async def test_empty_file_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_bytes(b"")
    result = await parse_file(path, file_id="e1")
    assert result.blocks == []
    assert result.aborted == "empty"
    assert any("文件为空" in w for w in result.warnings)


# ─────────────────── 类型清单 / 专用库 / 动态判定 ───────────────────

EXPECTED_EXTENSIONS = {
    ".pdf", ".docx", ".dotx", ".doc", ".wps", ".xlsx", ".xlsm", ".xls", ".et",
    ".pptx", ".odt", ".ods", ".odp", ".csv", ".tsv", ".json", ".xml", ".html",
    ".htm", ".rtf", ".txt", ".md", ".png", ".jpg",
}


def test_supported_extensions_are_the_single_source() -> None:
    """上传白名单由类型注册表派生，两者不允许各说各话。"""
    from app.modules.research.doc_gen import service

    extensions = set(supported_extensions())
    assert EXPECTED_EXTENSIONS <= extensions
    assert set(service.ALLOWED_EXTENSIONS) == extensions


def test_every_format_declares_a_library() -> None:
    """每种类型都必须写清使用的提取库，供运维核对与生成说明引用。"""
    formats = supported_formats()
    assert formats, "类型注册表不能为空"
    for item in formats:
        assert item["name"] and item["library"], item
        assert item["extensions"], item


async def test_format_chosen_by_magic_when_extension_missing(tmp_path: Path) -> None:
    docx = tmp_path / "real.docx"
    _make_docx(docx)
    disguised = tmp_path / "no_extension"
    disguised.write_bytes(docx.read_bytes())
    result = await parse_file(disguised, file_id="m1")
    assert result.format_name == "Word"
    assert any(b.kind == "heading" for b in result.blocks)


async def test_format_chosen_by_mime_hint(tmp_path: Path) -> None:
    path = tmp_path / "weird_name.dat"
    _make_xlsx(path)
    payload = path.read_bytes()
    target = tmp_path / "upload.dat"
    target.write_bytes(payload)
    result = await parse_file(
        target,
        file_id="mi1",
        mime_hint="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert result.format_name == "Excel"
    assert result.blocks


async def test_unsupported_container_is_not_salvaged_as_text(tmp_path: Path) -> None:
    """细分不出的 zip 容器不能把二进制碎片当正文——那会让模型产出看似合理的编造值。"""
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK\x03\x04" + os.urandom(4096))
    result = await parse_file(path, file_id="z2")
    assert result.blocks == []
    assert any("不支持的文件类型" in w for w in result.warnings)


async def test_random_binary_does_not_leak_as_garbage_text(tmp_path: Path) -> None:
    path = tmp_path / "noise.weird"
    path.write_bytes(os.urandom(8192))
    result = await parse_file(path, file_id="n1")
    assert result.blocks == []


async def test_legacy_doc_salvage(tmp_path: Path) -> None:
    text = "本文件描述了阿莫西林结晶工艺的研究过程，结晶温度为60℃。" * 6
    path = tmp_path / "legacy.doc"
    path.write_bytes(text.encode("utf-16-le"))
    result = await parse_file(path, file_id="d1")
    assert result.blocks
    assert any("抢救" in w for w in result.warnings)


# ─────────────────── 强制终止 ───────────────────


async def test_should_cancel_stops_parsing(tmp_path: Path) -> None:
    path = tmp_path / "long.txt"
    path.write_text("\n\n".join(f"第{i}段内容" for i in range(200)), encoding="utf-8")
    result = await parse_file(path, file_id="c1", should_cancel=lambda: True)
    assert result.aborted == "cancelled"
    assert any("强制终止" in w for w in result.warnings)


async def test_timeout_forces_abort_and_sets_event(tmp_path: Path, monkeypatch: Any) -> None:
    import threading
    import time as _time

    from app.modules.research.doc_gen import parsing as parsing_module

    def slow_sync(_path: Path, **_kwargs: Any) -> Any:
        _time.sleep(0.5)
        raise AssertionError("不应返回")

    monkeypatch.setattr(parsing_module, "parse_file_sync", slow_sync)
    path = tmp_path / "any.txt"
    path.write_text("内容", encoding="utf-8")
    event = threading.Event()
    result = await parse_file(path, file_id="t1", timeout_seconds=0.02, abort_event=event)
    assert result.aborted == "timeout"
    assert event.is_set(), "超时后必须通知线程内守卫退出"
    assert any("强制终止" in w for w in result.warnings)


async def test_truncation(tmp_path: Path) -> None:
    path = tmp_path / "long.txt"
    path.write_text("\n\n".join(["行内容" * 20] * 50), encoding="utf-8")
    result = await parse_file(path, file_id="long1", max_chars=300)

    assert result.char_count == 300
    assert any("文本超过上限已截断" in w for w in result.warnings)


async def test_image_ocr_success(tmp_path: Path, monkeypatch: Any) -> None:
    fake = FakeOCR()
    monkeypatch.setattr("app.shared.ocr_service.get_ocr_service", lambda: fake)
    path = tmp_path / "scan.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\nfake")

    result = await parse_file(path, file_id="img1")
    assert result.ocr_used is True
    assert [b.text for b in result.blocks] == ["结晶温度 60", "收率 92%"]
    assert len(fake.calls) == 1


async def test_image_ocr_unavailable(tmp_path: Path, monkeypatch: Any) -> None:
    def boom() -> Any:
        raise RuntimeError("OCR service not initialized")

    monkeypatch.setattr("app.shared.ocr_service.get_ocr_service", boom)
    path = tmp_path / "scan.jpg"
    path.write_bytes(b"\xff\xd8\xff\xe0fake")

    result = await parse_file(path, file_id="img2")
    assert result.blocks == []
    assert result.ocr_used is False
    assert any("OCR 服务不可用" in w for w in result.warnings)


def _make_blank_pdf(path: Path, pages: int = 1) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        writer.write(f)


async def test_pdf_scanned_ocr_fallback(tmp_path: Path, monkeypatch: Any) -> None:
    fake = FakeOCR(text="扫描件第一页文字")
    monkeypatch.setattr("app.shared.ocr_service.get_ocr_service", lambda: fake)
    monkeypatch.setattr("pdf2image.convert_from_path", lambda *a, **k: [object()])
    path = tmp_path / "scan.pdf"
    _make_blank_pdf(path, pages=1)

    result = await parse_file(path, file_id="pdf1")
    assert result.ocr_used is True
    assert [b.text for b in result.blocks] == ["扫描件第一页文字"]
    assert result.blocks[0].page == 1
    assert result.page_count == 1
    assert any("扫描件无文字层，已使用 OCR" in w for w in result.warnings)


async def test_pdf_all_empty_without_ocr_service(tmp_path: Path, monkeypatch: Any) -> None:
    def boom() -> Any:
        raise RuntimeError("OCR service not initialized")

    monkeypatch.setattr("app.shared.ocr_service.get_ocr_service", boom)
    path = tmp_path / "scan.pdf"
    _make_blank_pdf(path, pages=1)

    result = await parse_file(path, file_id="pdf2")
    assert result.blocks == []
    assert result.ocr_used is False
    assert any("OCR 服务不可用" in w for w in result.warnings)
    assert any("未能提取到文字" in w for w in result.warnings)
